import os
import re
import time
from datetime import datetime
from html import unescape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def replace_separator(words: str) -> str:
    """replace the "" separator with " & " """
    return words.replace("", " & ")


def rank_rules(rules: str) -> str:
    """rank the rules provided by parameters with 1.2.3...sign"""
    if len(rules.split(",")) >= 2:
        processed_rules = ""
        for rule_index in range(0, len(rules.split(","))):
            processed_rules += f'{rule_index+1}.{rules.split(",")[rule_index]}\n'
        return processed_rules.rstrip()
    else:
        return rules


def apply_paused(rules: str, paused_states: list) -> str:
    """Append (🚫Paused) to each rule whose paused flag equals "1".

    The paused list may carry one extra leading element (the "Input" slot
    which is not an actual rule), so both alignments are handled.
    """
    if not rules or not paused_states:
        return rules

    rule_lines = rules.split("\n")
    offset = len(paused_states) - len(rule_lines)
    marked = []

    for idx, rule in enumerate(rule_lines):
        if offset in (0, 1) and (idx + offset) < len(paused_states):
            flag = paused_states[idx + offset]
        else:
            flag = "0"

        marked.append(f"{rule}(🚫Paused)" if flag == "1" else rule)

    return "\n".join(marked)


class PageSplitter:
    """Split a slimedit page HTML into the import, intermediate and export
    sections using the <td> class markers.
    """

    IMPORT_MARKER = "source-columns"
    INTER_MARKER = "intermediate-columns"
    EXPORT_MARKER = "export-columns"

    @staticmethod
    def _find_td_start(html: str, marker: str) -> int:
        """Return the index of the '<td' whose class contains the marker."""
        pos = html.find(marker)

        if pos == -1:
            return -1

        td_pos = html.rfind("<td", 0, pos)

        return td_pos if td_pos != -1 else pos

    def split(self, html: str):
        """Return (import_section, inter_section, export_section, has_export).

        When the export section is missing, it is returned as an empty
        string and has_export is False.
        """
        idx_import = self._find_td_start(
            html,
            self.IMPORT_MARKER,
        )
        idx_inter = self._find_td_start(
            html,
            self.INTER_MARKER,
        )
        idx_export = self._find_td_start(
            html,
            self.EXPORT_MARKER,
        )

        if idx_import == -1 or idx_inter == -1:
            raise ValueError(
                "HTML does not contain the import / intermediate sections "
                "(source-columns / intermediate-columns)."
            )

        import_section = html[idx_import:idx_inter]

        if idx_export == -1:
            return import_section, html[idx_inter:], "", False

        return (
            import_section,
            html[idx_inter:idx_export],
            html[idx_export:],
            True,
        )


class ImportParser:
    """Extract attributes, connections and data source info from the
    import section of a slimedit page.
    """

    def __init__(self):
        self.field_pattern = r'(?<=field=").+?(?=")'
        self.attr_pattern = r'(?<=data-name=").+?(?=")'
        self.connection_pattern = r'(?<=connection=")[^"]+(?=" field=")'
        self.item_pattern = r'(?<=<li class="list-field source-field).+(?=>)'
        self.adit_suc_pattern = (
            r"(?<=</li><li class=).+" r'|(?<=data-original-title=").+(?="></i>)'
        )
        self.warning_pattern = r"&lt;strong&gt;Warning&lt;/strong&gt"

    def read(self, section: str) -> list:
        """Return [(attr_name, connection_str), ...].

        Import attributes without a connection are retained with a blank
        connection so they can still be included in the final report.
        """
        items = re.findall(self.item_pattern, section)
        result = []

        for item in items:
            attr_match = re.search(
                self.attr_pattern,
                item,
            )

            if not attr_match:
                continue

            conn_match = re.search(
                self.connection_pattern,
                item,
            )

            connection = replace_separator(conn_match.group(0)) if conn_match else " "

            result.append(
                (
                    attr_match.group(0),
                    connection,
                )
            )

        return result

    def read_additional(self, section: str) -> list:
        """Return [(attr_name, source_note), ...] for additional sources."""
        adit_suc = re.findall(
            self.adit_suc_pattern,
            section,
        )
        total = len(adit_suc)
        adis_pattern = r"(?<=Additional - ).+"

        paired = [
            (adit_suc[i], adit_suc[i + 1])
            for i in range(0, total - 1)
            if "Additional - " in adit_suc[i + 1]
        ]

        return [
            (
                re.search(
                    self.attr_pattern,
                    data[0],
                ).group(0),
                re.search(
                    adis_pattern,
                    data[1],
                ).group(0),
            )
            for data in paired
        ]

    def read_missing(self, section: str) -> list:
        """Return attr names whose import cell shows a Warning marker."""
        items = re.findall(
            self.item_pattern,
            section,
        )

        return [
            re.search(
                self.attr_pattern,
                item,
            ).group(0)
            for item in items
            if self.warning_pattern in item
        ]


class IntermediateParser:
    """Extract attributes and rules from the intermediate section."""

    def __init__(self):
        self.attr_pattern = r'(?<=search=").+?(?=")'
        self.field_pattern = r'(?<=field=").+?(?=")'
        self.connection_pattern = r'(?<=connection=")[^"]+(?=")'
        self.rbox_pattern = r'(?<=boxes-orig=")[^"]+(?=")'
        self.pause_pattern = r'(?<=boxes-paused=").+' r'(?=" boxes-paused-orig=)'
        self.item_pattern = r'<li class="list-field intermediate-field.*?</li>'

    def read(self, section: str) -> list:
        """Return (attr, field, connections, rules, paused) tuples.

        The connection attribute is read from whichever element carries a
        non-empty value. In the current platform build this is the trailing
        "connection-output output-intermediate" div; the "connection-input"
        div leaves the attribute empty. Field is always captured so an
        intermediate without a connection can still be matched by its id.
        """
        items = re.findall(
            self.item_pattern,
            section,
            flags=re.DOTALL,
        )
        result = []

        for item in items:
            attr_match = re.search(
                self.attr_pattern,
                item,
            )

            if not attr_match:
                continue

            attr_name = attr_match.group(0)

            field_match = re.search(
                self.field_pattern,
                item,
            )
            field_value = field_match.group(0) if field_match else " "

            conn_match = re.search(
                self.connection_pattern,
                item,
            )

            if not conn_match:
                result.append(
                    (
                        attr_name,
                        field_value,
                        " ",
                        " ",
                        [],
                    )
                )
                continue

            rbox_match = re.search(
                self.rbox_pattern,
                item,
            )
            pause_match = re.search(
                self.pause_pattern,
                item,
            )

            rules = ""

            if rbox_match:
                rules = rank_rules(
                    re.sub(
                        r"(Input,|Input)",
                        "",
                        rbox_match.group(0),
                    )
                )

            paused = pause_match.group(0).split(",") if pause_match else []

            result.append(
                (
                    attr_name,
                    field_value,
                    replace_separator(conn_match.group(0)),
                    rules,
                    paused,
                )
            )

        return result


class ExportParser:
    """Extract attributes and rules from the export section."""

    def __init__(self):
        self.attr_pattern = r'(?<=data-name=").+?(?=")'
        self.field_pattern = r'(?<=field=").+?(?=")'
        self.connection_pattern = r'(?<=connection=")[^"]+(?=" search=)'
        self.rbox_pattern = r'(?<=boxes-orig=")[^"]+(?=")'
        self.pause_pattern = r'(?<=boxes-paused=").+' r'(?=" boxes-paused-orig=)'
        self.item_pattern = r'<li class="list-field export-field.*?</li>'

    def read(self, section: str) -> list:
        """Return (attr, field, connections, rules, paused) tuples.

        Exports without a connection are kept with connection " " so they
        can be reported as standalone entries at the end of the report.
        """
        items = re.findall(
            self.item_pattern,
            section,
            flags=re.DOTALL,
        )
        result = []

        for item in items:
            attr_match = re.search(
                self.attr_pattern,
                item,
            )
            field_match = re.search(
                self.field_pattern,
                item,
            )

            if not (attr_match and field_match):
                continue

            conn_match = re.search(
                self.connection_pattern,
                item,
            )
            rbox_match = re.search(
                self.rbox_pattern,
                item,
            )
            pause_match = re.search(
                self.pause_pattern,
                item,
            )

            rules = ""

            if rbox_match:
                rules = rank_rules(
                    re.sub(
                        r"(Input,|Input)",
                        "",
                        rbox_match.group(0),
                    )
                )

            paused = pause_match.group(0).split(",") if pause_match else []

            conn_value = replace_separator(conn_match.group(0)) if conn_match else " "

            result.append(
                (
                    attr_match.group(0),
                    field_match.group(0),
                    conn_value,
                    rules,
                    paused,
                )
            )

        return result


class AttrInvestigator:
    """Can automatically extract attributes in import, intermediate and
    export from a single slimedit page HTML.
    """

    def __init__(self):
        self.import_parser = ImportParser()
        self.inter_parser = IntermediateParser()
        self.export_parser = ExportParser()

        self.page_file_name = ""
        self.report_file_name = ""
        self.has_export = False

        self.import_section = ""
        self.inter_section = ""
        self.export_section = ""
        self.complete_attrs = []

        self.sheet_header_4 = [
            "Import",
            "Rule",
            "Intermediate",
            "Data Source",
        ]
        self.sheet_header_6 = [
            "Import",
            "Rule",
            "Intermediate",
            "Export Rule",
            "Export",
            "Data Source",
        ]

        self.purple_fill = PatternFill(
            start_color="7030A0",
            end_color="7030A0",
            fill_type="solid",
        )
        self.green_fill = PatternFill(
            start_color="92D050",
            end_color="92D050",
            fill_type="solid",
        )
        self.blue_fill = PatternFill(
            start_color="00B0F0",
            end_color="00B0F0",
            fill_type="solid",
        )
        self.font_color = Font(
            color="FFFFFF",
        )

    @staticmethod
    def sanitize_file_name(file_name: str) -> str:
        """Return a safe file name without an extension."""
        file_name = unescape(file_name)

        file_name = re.sub(
            r'[<>:"/\\|?*\x00-\x1f]',
            "_",
            file_name,
        )

        file_name = re.sub(
            r"\s+",
            " ",
            file_name,
        ).strip()

        file_name = file_name.rstrip(". ")

        file_name = re.sub(
            r"\.(xlsx|xls|txt|html?)$",
            "",
            file_name,
            flags=re.IGNORECASE,
        ).strip()

        return file_name[:180]

    def extract_report_file_name(self, html: str) -> str:
        """Extract the report name from the export stage header.

        Raw HTML is checked first. If no matching header is found, the HTML
        is decoded and checked again to support entity-escaped HTML.
        """

        def extract_from_content(content: str) -> str:
            th_pattern = r"<th\b[^>]*>"

            for match in re.finditer(
                th_pattern,
                content,
                flags=re.IGNORECASE | re.DOTALL,
            ):
                opening_tag = match.group(0)

                class_match = re.search(
                    r'\bclass\s*=\s*(["\'])(.*?)\1',
                    opening_tag,
                    flags=re.IGNORECASE | re.DOTALL,
                )

                if not class_match:
                    continue

                class_names = class_match.group(2).split()

                if "export-th" not in class_names or "stage-name-th" not in class_names:
                    continue

                title_match = re.search(
                    r'\btitle\s*=\s*(["\'])(.*?)\1',
                    opening_tag,
                    flags=re.IGNORECASE | re.DOTALL,
                )

                if not title_match:
                    continue

                file_name = self.sanitize_file_name(title_match.group(2))

                if file_name:
                    return file_name

            return ""

        # First try the original HTML.
        file_name = extract_from_content(html)

        if file_name:
            return file_name

        # Fall back to decoded HTML for content containing &lt;th ... &gt;.
        decoded_html = unescape(html)

        if decoded_html != html:
            return extract_from_content(decoded_html)

        return ""

    def get_cur_dir(self):
        """List candidate text files and let the user pick one."""
        text_files = [txt for txt in os.listdir() if txt.endswith(".txt")]

        if not text_files:
            raise FileNotFoundError("No .txt file found in the current directory.")

        for index, text_file in enumerate(text_files):
            print(f"{index+1} => {text_file}")

        active = True

        while active:
            try:
                user_input = input(
                    "Type the number of the slimedit page " "you want to process> "
                )

                index = int(user_input.strip()) - 1

                if index < 0 or index >= len(text_files):
                    print("Please type a number within the list above")
                    continue

                self.page_file_name = text_files[index]
                active = False

                print(f"You have selected: {self.page_file_name}")

            except ValueError:
                print("Please type a number!")

    def load_page(self):
        """Read the HTML file and split it into the three sections."""
        html = Path(self.page_file_name).read_text(encoding="utf-8")

        # Prioritize and sanitize the user-selected TXT file name.
        selected_file_name = Path(self.page_file_name).stem
        report_name = self.sanitize_file_name(selected_file_name)

        # Fall back to the file name extracted from the HTML.
        if not report_name:
            report_name = self.extract_report_file_name(html)

        if report_name:
            timestamp = datetime.now().strftime("%Y%m%d%H%M")
            self.report_file_name = f"{report_name}_{timestamp}_mapping"
        else:
            self.report_file_name = ""

        splitter = PageSplitter()
        (
            self.import_section,
            self.inter_section,
            self.export_section,
            self.has_export,
        ) = splitter.split(html)

    def merge_all(self) -> list:
        """Combine import, intermediate and (optional) export into rows.

        Each import -> intermediate -> export chain becomes one row.
        Intermediate entries without a downstream export keep their row
        with the export columns blank. Export entries without an upstream
        intermediate are appended with the earlier columns blank.

        Matching rules:
          import.connection[] contains intermediate.field
          export.connection[] contains intermediate.field
        """
        import_data = self.import_parser.read(self.import_section)
        inter_data = self.inter_parser.read(self.inter_section)
        export_data = (
            self.export_parser.read(self.export_section) if self.has_export else []
        )

        processed_inter = []

        for attr, field, conn, rules, paused in inter_data:
            if rules and paused:
                rules = apply_paused(
                    rules,
                    paused,
                )

            processed_inter.append(
                (
                    attr,
                    field,
                    conn,
                    rules,
                )
            )

        processed_export = []

        for attr, field, conn, rules, paused in export_data:
            if rules and paused:
                rules = apply_paused(
                    rules,
                    paused,
                )

            processed_export.append(
                (
                    attr,
                    field,
                    conn,
                    rules,
                )
            )

        rows = []
        used_inter_attrs = set()
        used_export_attrs = set()

        for import_attr, import_conn in import_data:
            import_conn_set = set(import_conn.split(" & "))

            matched_inter = [
                entry for entry in processed_inter if entry[1] in import_conn_set
            ]

            if not matched_inter:
                rows.append(
                    [
                        import_attr,
                        " ",
                        " ",
                        " ",
                        " ",
                        " ",
                    ]
                    if self.has_export
                    else [
                        import_attr,
                        " ",
                        " ",
                        " ",
                    ]
                )
                continue

            for (
                inter_attr,
                inter_field,
                inter_conn,
                inter_rules,
            ) in matched_inter:
                used_inter_attrs.add(inter_attr)

                matched_export = [
                    entry
                    for entry in processed_export
                    if (
                        entry[2]
                        and inter_field
                        and inter_field != " "
                        and inter_field in entry[2].split(" & ")
                    )
                ]

                if not matched_export:
                    rows.append(
                        [
                            import_attr,
                            inter_rules,
                            inter_attr,
                            " ",
                            " ",
                            " ",
                        ]
                        if self.has_export
                        else [
                            import_attr,
                            inter_rules,
                            inter_attr,
                            " ",
                        ]
                    )
                    continue

                for (
                    exp_attr,
                    exp_field,
                    exp_conn,
                    exp_rules,
                ) in matched_export:
                    used_export_attrs.add(exp_attr)

                    rows.append(
                        [
                            import_attr,
                            inter_rules,
                            inter_attr,
                            exp_rules,
                            exp_attr,
                            " ",
                        ]
                        if self.has_export
                        else [
                            import_attr,
                            inter_rules,
                            inter_attr,
                            " ",
                        ]
                    )

        for attr, field, conn, rules in processed_inter:
            if attr in used_inter_attrs:
                continue

            matched_export = [
                entry
                for entry in processed_export
                if (
                    entry[2]
                    and field
                    and field != " "
                    and field in entry[2].split(" & ")
                )
            ]

            if not matched_export:
                rows.append(
                    [
                        " ",
                        " ",
                        attr,
                        " ",
                        " ",
                        " ",
                    ]
                    if self.has_export
                    else [
                        " ",
                        " ",
                        attr,
                        " ",
                    ]
                )
                continue

            for (
                exp_attr,
                exp_field,
                exp_conn,
                exp_rules,
            ) in matched_export:
                used_export_attrs.add(exp_attr)

                rows.append(
                    [
                        " ",
                        " ",
                        attr,
                        exp_rules,
                        exp_attr,
                        " ",
                    ]
                    if self.has_export
                    else [
                        " ",
                        " ",
                        attr,
                        " ",
                    ]
                )

        if self.has_export:
            for attr, field, conn, rules in processed_export:
                if attr in used_export_attrs:
                    continue

                rows.append(
                    [
                        " ",
                        " ",
                        " ",
                        rules,
                        attr,
                        " ",
                    ]
                )

        return rows

    def filter_source(self):
        """Assign Main/Additional/Missing source notes to each import row."""
        additional = self.import_parser.read_additional(self.import_section)
        missing = self.import_parser.read_missing(self.import_section)

        for row in self.complete_attrs:
            import_name = row[0]

            if import_name == " ":
                continue

            source = None

            for attr, note in additional:
                if attr == import_name:
                    source = note
                    break

            if source is None:
                source = "Missing Source" if import_name in missing else "Main Source"

            row[-1] = source

    def clear_data(self):
        """Clear the nonsense html character in the import and intermediate"""
        for complete_attr in self.complete_attrs:
            complete_attr[0] = re.sub(
                "(&gt;)",
                ">",
                re.sub(
                    "(&amp;)",
                    "&",
                    complete_attr[0],
                ),
            )

            complete_attr[2] = re.sub(
                "(&gt;)",
                ">",
                re.sub(
                    "(&amp;)",
                    "&",
                    complete_attr[2],
                ),
            )

            if self.has_export:
                complete_attr[4] = re.sub(
                    "(&gt;)",
                    ">",
                    re.sub(
                        "(&amp;)",
                        "&",
                        complete_attr[4],
                    ),
                )

            if len(complete_attr) >= 4:
                complete_attr[-1] = re.sub(
                    "&amp;amp;",
                    "&",
                    re.sub(
                        "&lt;br/&gt;Additional - ",
                        ";",
                        complete_attr[-1],
                    ),
                )

    def save_file(self):
        """Save the complete attributes to an Excel file with the format
        defined in format_excel.
        """
        wb = Workbook()
        sheet = wb.active
        sheet.title = "all_attrs_inter_export"

        sheet.append(self.sheet_header_6 if self.has_export else self.sheet_header_4)

        for row in self.complete_attrs:
            sheet.append(row)

        self.format_excel(sheet)

        if self.report_file_name:
            saved_file_name = f"{self.report_file_name}.xlsx"
        else:
            date_now = datetime.now().strftime("%Y_%m_%d_%H_%m_%S")
            saved_file_name = f"import_intermediate_attribute_" f"{date_now}.xlsx"

        wb.save(saved_file_name)

        print(f"Your file has been saved as {saved_file_name}")

    def format_excel(self, sheet):
        """Highlighting the unused attributes"""
        if self.has_export:
            main_columns = [
                "A",
                "B",
                "C",
                "D",
                "E",
                "F",
            ]
            headers = [
                "A1",
                "C1",
                "E1",
            ]
        else:
            main_columns = [
                "A",
                "B",
                "C",
                "D",
            ]
            headers = [
                "A1",
                "C1",
            ]

        for header in headers:
            sheet[f"{header}"].font = Font(
                color="FFFFFF",
                bold=True,
            )

        sheet["A1"].fill = self.green_fill
        sheet["C1"].fill = self.blue_fill

        if self.has_export:
            sheet["E1"].fill = self.green_fill

        for column in main_columns:
            sheet.column_dimensions[f"{column}"].width = 36

            sheet[f"{column}1"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for index, values in enumerate(
            self.complete_attrs,
            start=2,
        ):
            sheet[f"A{index}"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            sheet[f"C{index}"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            if self.has_export:
                sheet[f"E{index}"].alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )
                last_col = "F"
            else:
                last_col = "D"

            sheet[f"{last_col}{index}"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            if values[0] == " " and values[2] != " ":
                sheet[f"C{index}"].fill = self.purple_fill
                sheet[f"C{index}"].font = self.font_color

            if values[0] != " " and values[2] == " ":
                sheet[f"A{index}"].fill = self.purple_fill
                sheet[f"A{index}"].font = self.font_color

    def start_process(self) -> None:
        """Starting to read html file and processing data"""
        start = time.time()

        self.get_cur_dir()
        self.load_page()

        self.complete_attrs = self.merge_all()

        self.filter_source()
        self.clear_data()

        end = time.time()

        print(f"Total processing time: " f"{round(end-start, 2)}s")

        return None


if __name__ == "__main__":
    attr_inves = AttrInvestigator()
    attr_inves.start_process()
    attr_inves.save_file()
